"""Import de produits depuis un fichier Excel (.xlsx) et modèle à télécharger."""

import re
import unicodedata
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook

MAX_LIGNES = 1000

COLONNES_MODELE = [
    ("Nom", "nom"),
    ("Unité", "unite"),
    ("Prix de vente", "prix_vente"),
    ("Prix d'achat", "cout_initial"),
    ("Stock initial", "stock_initial"),
    ("Seuil d'alerte", "seuil_alerte"),
    ("Code article", "sku"),
]

# En-têtes reconnus (après suppression des accents, ponctuation et majuscules).
SYNONYMES = {
    "nom": {"nom", "produit", "article", "designation", "libelle", "nom du produit"},
    "unite": {"unite", "unite de mesure"},
    "prix_vente": {"prix de vente", "prix vente", "pv", "prix"},
    "cout_initial": {
        "prix d achat", "prix achat", "cout", "cout achat", "cout d achat",
        "cout unitaire", "pa",
    },
    "stock_initial": {"stock initial", "stock", "quantite", "qte", "stock actuel"},
    "seuil_alerte": {"seuil d alerte", "seuil alerte", "seuil", "alerte"},
    "sku": {"code article", "code", "sku", "reference", "ref"},
}
CHAMPS_NOMBRES = {"prix_vente", "cout_initial", "stock_initial", "seuil_alerte"}


class ImportInvalideError(ValueError):
    """Fichier inexploitable (format, colonnes, taille)."""


def _normaliser(texte) -> str:
    sans_accents = (
        unicodedata.normalize("NFKD", str(texte or ""))
        .encode("ascii", "ignore")
        .decode()
        .lower()
    )
    return re.sub(r"[^a-z0-9]+", " ", sans_accents).strip()


def _nombre(valeur):
    """Convertit une cellule en texte décimal (ou None si vide). Lève ValueError."""
    if valeur is None:
        return None
    if isinstance(valeur, bool):
        raise ValueError("nombre attendu")
    if isinstance(valeur, (int, float, Decimal)):
        return format(Decimal(str(valeur)), "f")
    texte = str(valeur).replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    if not texte:
        return None
    try:
        return format(Decimal(texte), "f")
    except InvalidOperation:
        raise ValueError(f"« {valeur} » n'est pas un nombre")


def lire_fichier_produits(fichier) -> list:
    """Lit le fichier et renvoie [{"ligne", "donnees", "erreur"}].

    `ligne` est le numéro de ligne dans Excel (l'en-tête est la ligne 1 en
    général). Une ligne illisible a `erreur` renseignée et pas de `donnees`.
    """
    try:
        classeur = load_workbook(fichier, read_only=True, data_only=True)
    except Exception:
        raise ImportInvalideError("Fichier illisible : envoyez un fichier Excel .xlsx.")

    feuille = classeur["Produits"] if "Produits" in classeur.sheetnames else classeur.worksheets[0]

    colonnes = None  # champ -> index de colonne
    resultat = []
    for numero, cellules in enumerate(feuille.iter_rows(values_only=True), start=1):
        if colonnes is None:
            if not any(c not in (None, "") for c in cellules):
                continue  # lignes vides avant l'en-tête
            colonnes = {}
            for index, titre in enumerate(cellules):
                cle = _normaliser(titre)
                for champ, noms in SYNONYMES.items():
                    if cle in noms and champ not in colonnes:
                        colonnes[champ] = index
            if "nom" not in colonnes:
                raise ImportInvalideError(
                    "Colonne « Nom » introuvable. Utilisez le modèle Excel proposé."
                )
            continue

        if not any(c not in (None, "") for c in cellules):
            continue  # ligne vide
        if len(resultat) >= MAX_LIGNES:
            raise ImportInvalideError(f"Maximum {MAX_LIGNES} produits par fichier.")

        donnees, erreur = {}, None
        for champ, index in colonnes.items():
            valeur = cellules[index] if index < len(cellules) else None
            if champ in CHAMPS_NOMBRES:
                try:
                    nombre = _nombre(valeur)
                except ValueError as exc:
                    erreur = f"{dict((c, t) for t, c in COLONNES_MODELE)[champ]} : {exc}"
                    break
                if nombre is not None:
                    donnees[champ] = nombre
            elif valeur not in (None, ""):
                donnees[champ] = str(valeur).strip()
        resultat.append({"ligne": numero, "donnees": None if erreur else donnees, "erreur": erreur})
    return resultat


def modele_xlsx() -> bytes:
    """Fichier Excel modèle (en-têtes + feuille d'aide)."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Produits"
    feuille.append([titre for titre, _ in COLONNES_MODELE])
    for index, largeur in enumerate([30, 12, 16, 16, 16, 16, 16], start=1):
        feuille.column_dimensions[chr(64 + index)].width = largeur

    aide = classeur.create_sheet("Aide")
    for ligne in [
        ["Comment remplir le fichier"],
        [],
        ["Nom", "Obligatoire. Un produit par ligne, sans doublon."],
        ["Unité", "sac, pièce, litre, kg… (« pièce » par défaut)."],
        ["Prix de vente", "Prix de vente unitaire."],
        ["Prix d'achat", "Coût d'achat unitaire (sert au coût moyen du stock)."],
        ["Stock initial", "Quantité réellement en stock aujourd'hui."],
        ["Seuil d'alerte", "Femi vous prévient quand le stock atteint ce niveau."],
        ["Code article", "Facultatif, unique s'il est renseigné."],
        [],
        ["Les lignes déjà existantes (même nom) sont ignorées, jamais écrasées."],
    ]:
        aide.append(ligne)
    aide.column_dimensions["A"].width = 18
    aide.column_dimensions["B"].width = 70

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
