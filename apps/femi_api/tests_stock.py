from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from apps.femi_account.models import Entreprise, MouvementStock, Produit, Utilisateur


class StockAPITests(APITestCase):
    def setUp(self):
        self.entreprise = Entreprise.objects.create(nom="Boutique Test", devise="XOF")
        self.admin = Utilisateur.objects.create_user(
            username="admin1", email="admin1@test.com", password="x", entreprise=self.entreprise, role="ADMIN"
        )
        self.employe = Utilisateur.objects.create_user(
            username="emp1", email="emp1@test.com", password="x", entreprise=self.entreprise, role="EMPLOYE"
        )
        self.client.force_authenticate(self.admin)

    # -- utilitaires -------------------------------------------------
    def _produit(self, nom="Ciment", stock="10", cout="100", **extra):
        corps = {"nom": nom, "unite": "sac", "prix_vente": "150", **extra}
        if stock is not None:
            corps["stock_initial"] = stock
        if cout is not None:
            corps["cout_initial"] = cout
        r = self.client.post(reverse("api_stock_produits"), corps, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        return r.json()

    def _mouvement(self, produit_id, type_mouvement, quantite, **extra):
        return self.client.post(
            reverse("api_stock_mouvements"),
            {"produit": produit_id, "type_mouvement": type_mouvement, "quantite": quantite, **extra},
            format="json",
        )

    def _produit_en_base(self, produit_id):
        return Produit.objects.get(pk=produit_id)

    # -- tests -------------------------------------------------------
    def test_creation_avec_stock_initial_cree_un_mouvement(self):
        p = self._produit(stock="10", cout="100")
        self.assertEqual(Decimal(p["stock_actuel"]), Decimal("10"))
        self.assertEqual(Decimal(p["cout_moyen"]), Decimal("100"))
        m = MouvementStock.objects.get(produit_id=p["id"])
        self.assertEqual((m.type_mouvement, m.motif), ("ENTREE", "STOCK_INITIAL"))
        self.assertEqual(m.stock_apres, Decimal("10"))

    def test_entree_recalcule_le_cout_moyen_pondere(self):
        p = self._produit(stock="10", cout="100")
        r = self._mouvement(p["id"], "ENTREE", "10", cout_unitaire="200")
        self.assertEqual(r.status_code, 201, r.content)
        produit = self._produit_en_base(p["id"])
        self.assertEqual(produit.stock_actuel, Decimal("20"))
        self.assertEqual(produit.cout_moyen, Decimal("150.00"))

    def test_sortie_diminue_le_stock_et_fige_le_cout(self):
        p = self._produit(stock="10", cout="100")
        r = self._mouvement(p["id"], "SORTIE", "4")
        self.assertEqual(r.status_code, 201, r.content)
        corps = r.json()
        self.assertEqual(Decimal(corps["produit"]["stock_actuel"]), Decimal("6"))
        self.assertEqual(Decimal(corps["mouvement"]["variation"]), Decimal("-4"))
        self.assertEqual(Decimal(corps["mouvement"]["cout_unitaire"]), Decimal("100"))

    def test_sortie_au_dela_du_stock_est_refusee(self):
        p = self._produit(stock="3", cout="100")
        r = self._mouvement(p["id"], "SORTIE", "5")
        self.assertEqual(r.status_code, 409)
        self.assertEqual(r.json()["code"], "STOCK_INSUFFISANT")
        self.assertEqual(self._produit_en_base(p["id"]).stock_actuel, Decimal("3"))
        self.assertEqual(MouvementStock.objects.filter(produit_id=p["id"]).count(), 1)

    def test_idempotence_ne_compte_qu_une_fois(self):
        p = self._produit(stock="10", cout="100")
        r1 = self._mouvement(p["id"], "SORTIE", "2", cle_idempotence="abc-123")
        r2 = self._mouvement(p["id"], "SORTIE", "2", cle_idempotence="abc-123")
        self.assertEqual((r1.status_code, r2.status_code), (201, 200))
        self.assertEqual(r1.json()["mouvement"]["id"], r2.json()["mouvement"]["id"])
        self.assertEqual(self._produit_en_base(p["id"]).stock_actuel, Decimal("8"))

    def test_idempotence_par_en_tete_http(self):
        p = self._produit(stock="10", cout="100")
        for _ in range(2):
            self.client.post(
                reverse("api_stock_mouvements"),
                {"produit": p["id"], "type_mouvement": "SORTIE", "quantite": "1"},
                format="json",
                HTTP_IDEMPOTENCY_KEY="cle-header",
            )
        self.assertEqual(self._produit_en_base(p["id"]).stock_actuel, Decimal("9"))

    def test_ajustement_d_inventaire(self):
        p = self._produit(stock="10", cout="100")
        r = self._mouvement(p["id"], "AJUSTEMENT", "7")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Decimal(r.json()["mouvement"]["variation"]), Decimal("-3"))
        self.assertEqual(self._produit_en_base(p["id"]).stock_actuel, Decimal("7"))
        # même valeur que le stock actuel : refusé
        self.assertEqual(self._mouvement(p["id"], "AJUSTEMENT", "7").status_code, 400)

    def test_motif_incoherent_avec_le_type_est_refuse(self):
        p = self._produit()
        self.assertEqual(self._mouvement(p["id"], "ENTREE", "1", motif="VENTE").status_code, 400)

    def test_quantite_nulle_refusee(self):
        p = self._produit()
        self.assertEqual(self._mouvement(p["id"], "SORTIE", "0").status_code, 400)

    def test_employe_peut_vendre_mais_rien_d_autre(self):
        p = self._produit(stock="10", cout="100")
        self.client.force_authenticate(self.employe)
        self.assertEqual(self._mouvement(p["id"], "SORTIE", "1").status_code, 201)
        self.assertEqual(self._mouvement(p["id"], "SORTIE", "1", motif="PERTE").status_code, 403)
        self.assertEqual(self._mouvement(p["id"], "ENTREE", "5").status_code, 403)
        self.assertEqual(self._mouvement(p["id"], "AJUSTEMENT", "2").status_code, 403)
        r = self.client.post(reverse("api_stock_produits"), {"nom": "X"}, format="json")
        self.assertEqual(r.status_code, 403)
        r = self.client.patch(
            reverse("api_stock_produit_detail", args=[p["id"]]), {"prix_vente": "1"}, format="json"
        )
        self.assertEqual(r.status_code, 403)
        # mais il peut consulter
        self.assertEqual(self.client.get(reverse("api_stock_produits")).status_code, 200)

    def test_isolation_entre_entreprises(self):
        p = self._produit()
        autre = Entreprise.objects.create(nom="Autre", devise="XOF")
        intrus = Utilisateur.objects.create_user(
            username="autre1", email="autre1@test.com", password="x", entreprise=autre, role="ADMIN"
        )
        self.client.force_authenticate(intrus)
        self.assertEqual(self._mouvement(p["id"], "SORTIE", "1").status_code, 404)
        r = self.client.get(reverse("api_stock_produit_detail", args=[p["id"]]))
        self.assertEqual(r.status_code, 404)
        self.assertEqual(self.client.get(reverse("api_stock_produits")).json()["total"], 0)
        self.assertEqual(self.client.get(reverse("api_stock_mouvements")).json()["total"], 0)

    def test_le_journal_est_immuable(self):
        p = self._produit()
        m = MouvementStock.objects.get(produit_id=p["id"])
        m.note = "triche"
        with self.assertRaises(ValueError):
            m.save()
        with self.assertRaises(ValueError):
            m.delete()

    def test_un_produit_avec_historique_ne_peut_pas_etre_supprime(self):
        from django.db.models import ProtectedError

        p = self._produit()
        with self.assertRaises(ProtectedError):
            self._produit_en_base(p["id"]).delete()

    def test_nom_en_double_insensible_a_la_casse(self):
        self._produit(nom="Ciment")
        r = self.client.post(reverse("api_stock_produits"), {"nom": "ciment"}, format="json")
        self.assertEqual(r.status_code, 400)

    def test_sku_unique_mais_plusieurs_sku_vides_acceptes(self):
        self._produit(nom="A", sku="REF-1", stock=None, cout=None)
        self._produit(nom="B", sku="", stock=None, cout=None)
        self._produit(nom="C", sku="", stock=None, cout=None)
        r = self.client.post(
            reverse("api_stock_produits"), {"nom": "D", "sku": "ref-1"}, format="json"
        )
        self.assertEqual(r.status_code, 400)

    def test_le_stock_ne_se_modifie_pas_par_patch(self):
        p = self._produit()
        r = self.client.patch(
            reverse("api_stock_produit_detail", args=[p["id"]]),
            {"stock_initial": "999"},
            format="json",
        )
        self.assertEqual(r.status_code, 400)
        r = self.client.patch(
            reverse("api_stock_produit_detail", args=[p["id"]]),
            {"stock_actuel": "999", "prix_vente": "175"},
            format="json",
        )
        self.assertEqual(r.status_code, 200)
        produit = self._produit_en_base(p["id"])
        self.assertEqual(produit.stock_actuel, Decimal("10"))
        self.assertEqual(produit.prix_vente, Decimal("175"))

    def test_produit_archive_refuse_les_mouvements_et_disparait_de_la_liste(self):
        p = self._produit()
        r = self.client.patch(
            reverse("api_stock_produit_detail", args=[p["id"]]), {"actif": False}, format="json"
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self._mouvement(p["id"], "SORTIE", "1").status_code, 400)
        self.assertEqual(self.client.get(reverse("api_stock_produits")).json()["total"], 0)
        r = self.client.get(reverse("api_stock_produits") + "?actif=false")
        self.assertEqual(r.json()["total"], 1)

    def test_resume_alertes_ruptures_et_valeur(self):
        self._produit(nom="A", stock="10", cout="100", seuil_alerte="3")   # OK : 1000
        p_b = self._produit(nom="B", stock="2", cout="50", seuil_alerte="3")  # ALERTE : 100
        p_c = self._produit(nom="C", stock="1", cout="10")                  # sera en rupture
        self._mouvement(p_c["id"], "SORTIE", "1")
        r = self.client.get(reverse("api_stock_resume")).json()
        self.assertEqual(r["nombre_produits"], 3)
        self.assertEqual(r["nb_alertes"], 1)
        self.assertEqual(r["nb_ruptures"], 1)
        self.assertEqual(Decimal(r["valeur_stock"]), Decimal("1100.00"))
        # filtres de liste
        alerte = self.client.get(reverse("api_stock_produits") + "?statut=ALERTE").json()
        self.assertEqual([x["nom"] for x in alerte["resultats"]], ["B"])
        rupture = self.client.get(reverse("api_stock_produits") + "?statut=RUPTURE").json()
        self.assertEqual([x["nom"] for x in rupture["resultats"]], ["C"])
        self.assertEqual(p_b["nom"], "B")

    def test_historique_filtre_par_produit_et_pagination(self):
        a = self._produit(nom="A", stock="10", cout="1")
        b = self._produit(nom="B", stock="10", cout="1")
        for _ in range(3):
            self._mouvement(a["id"], "SORTIE", "1")
        self._mouvement(b["id"], "SORTIE", "1")
        r = self.client.get(reverse("api_stock_mouvements") + f"?produit={a['id']}").json()
        self.assertEqual(r["total"], 4)  # stock initial + 3 sorties
        r = self.client.get(reverse("api_stock_mouvements") + f"?produit={a['id']}&limit=2").json()
        self.assertEqual((r["total"], len(r["resultats"])), (4, 2))
        bad = self.client.get(reverse("api_stock_mouvements") + "?date_debut=pas-une-date")
        self.assertEqual(bad.status_code, 400)

    def test_quantites_decimales(self):
        p = self._produit(nom="Riz", stock="12.500", cout="500")
        r = self._mouvement(p["id"], "SORTIE", "2.250")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(self._produit_en_base(p["id"]).stock_actuel, Decimal("10.250"))


def _xlsx(lignes, titre="Produits"):
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = titre
    for ligne in lignes:
        feuille.append(ligne)
    tampon = BytesIO()
    classeur.save(tampon)
    return SimpleUploadedFile(
        "produits.xlsx",
        tampon.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


class ImportProduitsTests(APITestCase):
    def setUp(self):
        self.entreprise = Entreprise.objects.create(nom="Boutique Import", devise="XOF")
        self.admin = Utilisateur.objects.create_user(
            username="adm", email="adm@test.com", password="x",
            entreprise=self.entreprise, role="ADMIN",
        )
        self.client.force_authenticate(self.admin)
        self.url = reverse("api_stock_import")

    def _importer(self, lignes):
        return self.client.post(self.url, {"fichier": _xlsx(lignes)}, format="multipart")

    def test_import_cree_les_produits_et_le_stock_initial(self):
        r = self._importer([
            ["Nom", "Unité", "Prix de vente", "Prix d'achat", "Stock initial", "Seuil d'alerte", "Code article"],
            ["Ciment", "sac", 150, 100, 10, 3, "CIM-1"],
            ["Peinture", "litre", "2 500,50", "1 800", "4,5", None, None],
        ])
        self.assertEqual(r.status_code, 200, r.content)
        corps = r.json()
        self.assertEqual((corps["total_lignes"], corps["crees"], corps["ignores"]), (2, 2, []))
        ciment = Produit.objects.get(nom="Ciment")
        self.assertEqual((ciment.stock_actuel, ciment.cout_moyen, ciment.sku), (Decimal("10"), Decimal("100"), "CIM-1"))
        peinture = Produit.objects.get(nom="Peinture")
        self.assertEqual(peinture.stock_actuel, Decimal("4.5"))
        self.assertEqual(peinture.prix_vente, Decimal("2500.50"))
        self.assertEqual(MouvementStock.objects.filter(motif="STOCK_INITIAL").count(), 2)

    def test_lignes_en_erreur_signalees_sans_bloquer_les_autres(self):
        Produit.objects.create(entreprise=self.entreprise, nom="Existant")
        r = self._importer([
            ["Produit", "Quantité", "Prix"],
            ["Bon", 5, 100],
            ["Existant", 1, 1],
            ["Bon", 2, 2],
            ["Mauvais", "beaucoup", 10],
            [None, 3, 3],
        ])
        corps = r.json()
        self.assertEqual(corps["crees"], 1)
        erreurs = {i["ligne"]: i["erreur"] for i in corps["ignores"]}
        self.assertEqual(set(erreurs), {3, 4, 5, 6})
        self.assertIn("déjà", erreurs[3])
        self.assertIn("Doublon", erreurs[4])
        self.assertIn("nombre", erreurs[5])
        self.assertEqual(Produit.objects.filter(entreprise=self.entreprise).count(), 2)

    def test_en_tete_obligatoire(self):
        r = self._importer([["Foo", "Bar"], [1, 2]])
        self.assertEqual(r.status_code, 400)
        self.assertIn("Nom", r.json()["error"])

    def test_fichier_non_excel_refuse(self):
        faux = SimpleUploadedFile("x.xlsx", b"ce n'est pas un excel")
        r = self.client.post(self.url, {"fichier": faux}, format="multipart")
        self.assertEqual(r.status_code, 400)

    def test_fichier_manquant(self):
        self.assertEqual(self.client.post(self.url, {}, format="multipart").status_code, 400)

    def test_employe_ne_peut_pas_importer(self):
        emp = Utilisateur.objects.create_user(
            username="emp", email="emp@test.com", password="x",
            entreprise=self.entreprise, role="EMPLOYE",
        )
        self.client.force_authenticate(emp)
        self.assertEqual(self._importer([["Nom"], ["A"]]).status_code, 403)

    def test_le_modele_telecharge_est_reimportable(self):
        r = self.client.get(reverse("api_stock_modele"))
        self.assertEqual(r.status_code, 200)
        classeur = load_workbook(BytesIO(r.content))
        self.assertEqual(classeur.sheetnames, ["Produits", "Aide"])
        feuille = classeur["Produits"]
        feuille.append(["Sable", "sac", 100, 60, 8, 2, ""])
        tampon = BytesIO()
        classeur.save(tampon)
        fichier = SimpleUploadedFile("m.xlsx", tampon.getvalue())
        r = self.client.post(self.url, {"fichier": fichier}, format="multipart")
        self.assertEqual(r.json()["crees"], 1, r.content)
        self.assertEqual(Produit.objects.get(nom="Sable").stock_actuel, Decimal("8"))
